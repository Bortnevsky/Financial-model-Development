#!/usr/bin/env python3
"""
Конвертер Excel в JSON для Financial Model Assistant
"""

import sys
import json
from pathlib import Path
from datetime import datetime, date
import warnings
warnings.filterwarnings('ignore')

try:
    import openpyxl
except ImportError:
    print("Установите openpyxl: pip install openpyxl")
    sys.exit(1)


def convert_excel_to_json(excel_path: str, output_path: str = None) -> str:
    """
    Конвертирует Excel файл в JSON формат для анализа.

    Args:
        excel_path: путь к Excel файлу
        output_path: путь для сохранения JSON (опционально)

    Returns:
        путь к созданному JSON файлу
    """
    excel_path = Path(excel_path)

    if not excel_path.exists():
        raise FileNotFoundError(f"Файл не найден: {excel_path}")

    if not output_path:
        output_path = excel_path.with_suffix('.json')
    else:
        output_path = Path(output_path)

    print(f"Загрузка {excel_path}...")
    wb = openpyxl.load_workbook(str(excel_path), data_only=False)

    model = {}
    total_cells = 0

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        sheet_data = {
            "rows": sheet.max_row,
            "cols": sheet.max_column,
            "cells": {},
            "merged": []  # объединённые ячейки
        }

        # Сохраняем информацию об объединённых ячейках
        for merged_range in sheet.merged_cells.ranges:
            sheet_data["merged"].append({
                "range": str(merged_range),
                "start": merged_range.min_col,
                "end": merged_range.max_col,
                "top": merged_range.min_row,
                "bottom": merged_range.max_row
            })

        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell_addr = cell.coordinate
                    value = cell.value

                    cell_info = {}

                    # Если это формула
                    if isinstance(value, str) and value.startswith("="):
                        cell_info["f"] = value  # formula
                    elif isinstance(value, (datetime, date)):
                        cell_info["v"] = value.isoformat()
                        cell_info["t"] = "d"  # date type
                    elif isinstance(value, (int, float)):
                        cell_info["v"] = value
                    elif isinstance(value, bool):
                        cell_info["v"] = value
                    elif isinstance(value, str):
                        cell_info["v"] = value
                    else:
                        # Для всех остальных типов - конвертируем в строку
                        cell_info["v"] = str(value)
                        cell_info["t"] = "special"

                    if cell.number_format and cell.number_format != "General":
                        cell_info["fmt"] = cell.number_format

                    sheet_data["cells"][cell_addr] = cell_info
                    total_cells += 1

        model[sheet_name] = sheet_data
        cells_count = len(sheet_data["cells"])
        print(f"  {sheet_name}: {cells_count} ячеек")

    # Сохраняем
    print(f"\nСохранение в {output_path}...")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(model, f, ensure_ascii=False)

    size_mb = output_path.stat().st_size / 1024 / 1024
    print(f"\nГотово!")
    print(f"  Файл: {output_path}")
    print(f"  Размер: {size_mb:.2f} MB")
    print(f"  Листов: {len(model)}")
    print(f"  Всего ячеек: {total_cells}")

    return str(output_path)


def main():
    if len(sys.argv) < 2:
        # Ищем Excel файлы в текущей директории
        excel_files = list(Path(".").glob("*.xlsx")) + list(Path(".").glob("*.xls"))

        if not excel_files:
            print("Использование: python excel_to_json.py <файл.xlsx> [выходной.json]")
            print("\nExcel файлы не найдены в текущей директории")
            sys.exit(1)

        print("Доступные Excel файлы:")
        for i, f in enumerate(excel_files, 1):
            size_mb = f.stat().st_size / 1024 / 1024
            print(f"  {i}. {f.name} ({size_mb:.2f} MB)")

        choice = input("\nВыберите файл (номер): ").strip()

        try:
            idx = int(choice) - 1
            if 0 <= idx < len(excel_files):
                excel_path = str(excel_files[idx])
            else:
                print("Неверный номер")
                sys.exit(1)
        except ValueError:
            print("Введите номер файла")
            sys.exit(1)
    else:
        excel_path = sys.argv[1]

    output_path = sys.argv[2] if len(sys.argv) > 2 else None

    try:
        convert_excel_to_json(excel_path, output_path)
    except Exception as e:
        print(f"Ошибка: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
